"""One-off generator for the sample Excel/PDF files used as connector demo
data for the Finance and Railways ministries. Run with:

    python scripts/generate_samples.py

Re-run any time you want to regenerate/reset the sample source files.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from fpdf import FPDF, XPos, YPos

ROOT = Path(__file__).resolve().parent.parent


def write_excel(path: Path, sheet_name: str, header: list[str], rows: list[list]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(header)
    for row in rows:
        ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"wrote {path}")


def write_pdf_table(path: Path, title: str, header: list[str], rows: list[list[str]]) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)

    col_width = 90
    row_height = 8

    pdf.set_font("Helvetica", "B", 11)
    for col in header:
        pdf.cell(col_width, row_height, col, border=1)
    pdf.ln(row_height)

    pdf.set_font("Helvetica", "", 11)
    for row in rows:
        for cell in row:
            pdf.cell(col_width, row_height, str(cell), border=1)
        pdf.ln(row_height)

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    print(f"wrote {path}")


def main() -> None:
    # Finance
    write_excel(
        ROOT / "data/finance/revenue_report.xlsx",
        sheet_name="Revenue",
        header=["Head", "Amount (Cr)", "Period"],
        rows=[
            ["Gross Tax Revenue", 2534567, "FY 2025-26"],
            ["Corporation Tax", 987654, "FY 2025-26"],
            ["Income Tax", 876543, "FY 2025-26"],
            ["GST Collections", 670370, "FY 2025-26"],
        ],
    )
    write_pdf_table(
        ROOT / "data/finance/budget_report.pdf",
        title="Union Budget - Expenditure Summary (FY 2025-26)",
        header=["Item", "Amount (Rs Cr)"],
        rows=[
            ["Capital Expenditure", "11,12,650"],
            ["Revenue Expenditure", "37,28,920"],
            ["Total Expenditure", "48,41,570"],
        ],
    )

    # Railways
    write_excel(
        ROOT / "data/railways/freight_report.xlsx",
        sheet_name="Freight",
        header=["Head", "Amount (Million Tonnes)", "Period"],
        rows=[
            ["Freight Loading", 1618.5, "FY 2025-26"],
            ["Coal Freight", 785.2, "FY 2025-26"],
            ["Container Freight", 102.4, "FY 2025-26"],
        ],
    )
    write_pdf_table(
        ROOT / "data/railways/performance_report.pdf",
        title="Indian Railways - Performance Report (FY 2025-26)",
        header=["Item", "Value (%)"],
        rows=[
            ["On-Time Performance", "82.6"],
            ["Punctuality Improvement", "4.1"],
            ["Safety Compliance Rate", "99.2"],
        ],
    )


if __name__ == "__main__":
    main()
